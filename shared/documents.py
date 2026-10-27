from pathlib import Path
import html

def read_text(path):
    path=Path(path)
    if path.stat().st_size>100*1024**2:raise ValueError('Document trop volumineux.')
    if path.suffix.lower()=='.pdf':
        import pymupdf
        with pymupdf.open(path) as document:
            if document.needs_pass:raise ValueError('PDF protégé par mot de passe.')
            return '\n\n'.join(page.get_text(sort=True) for page in document)
    if path.suffix.lower()=='.docx':
        from docx import Document
        document=Document(path)
        return '\n'.join(p.text for p in document.paragraphs)
    return path.read_text(encoding='utf-8-sig')

def save_docx(path,text,font='Arial',subject=''):
    from docx import Document
    from docx.shared import Pt
    document=Document()
    style=document.styles['Normal'];style.font.name=font;style.font.size=Pt(14)
    if subject:document.sections[0].header.paragraphs[0].text=subject
    for line in text.split('\n'):document.add_paragraph(line)
    document.save(path)

def save_rich_docx(path,qt_document,subject=''):
    from docx import Document
    from docx.shared import Pt
    document=Document()
    if subject:document.sections[0].header.paragraphs[0].text=subject
    block=qt_document.begin()
    while block.isValid():
        paragraph=document.add_paragraph();iterator=block.begin()
        while not iterator.atEnd():
            fragment=iterator.fragment()
            if fragment.isValid():
                run=paragraph.add_run(fragment.text());style=fragment.charFormat();run.bold=style.fontWeight()>=700;run.italic=style.fontItalic();run.underline=style.fontUnderline()
                run.font.name=style.font().family() or 'Arial';run.font.size=Pt(style.fontPointSize() or 14)
            iterator+=1
        block=block.next()
    document.save(path)
